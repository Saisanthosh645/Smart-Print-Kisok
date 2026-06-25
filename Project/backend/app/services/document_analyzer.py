import io
import re
from typing import Any

from docx import Document as DocxDocument
from pptx import Presentation
from PyPDF2 import PdfReader

from app.models.document import FileType


def parse_page_range(page_range: str | None, total_pages: int) -> list[int]:
    if not page_range or page_range.strip().lower() in ("all", ""):
        return list(range(1, total_pages + 1))

    pages: set[int] = set()
    for part in page_range.split(","):
        part = part.strip()
        if "-" in part:
            start_s, end_s = part.split("-", 1)
            start, end = int(start_s), int(end_s)
            pages.update(range(start, end + 1))
        else:
            pages.add(int(part))
    return sorted(p for p in pages if 1 <= p <= total_pages)


def count_pages(content: bytes, file_type: FileType) -> int:
    if file_type == FileType.PDF:
        reader = PdfReader(io.BytesIO(content))
        return len(reader.pages)
    if file_type == FileType.DOCX:
        doc = DocxDocument(io.BytesIO(content))
        # Approximate: ~300 words per page
        words = sum(len(p.text.split()) for p in doc.paragraphs)
        return max(1, words // 300 + (1 if words % 300 else 0))
    if file_type == FileType.PPTX:
        prs = Presentation(io.BytesIO(content))
        return len(prs.slides)
    return 1


def analyze_document(content: bytes, file_type: FileType) -> dict[str, Any]:
    """AI-style document analysis: blank pages, duplicates, orientation."""
    analysis: dict[str, Any] = {
        "blank_pages": [],
        "duplicate_pages": [],
        "orientation_issues": [],
        "recommendations": [],
    }

    if file_type != FileType.PDF:
        analysis["recommendations"].append("Convert to PDF for best print quality.")
        return analysis

    reader = PdfReader(io.BytesIO(content))
    page_hashes: dict[str, list[int]] = {}

    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if len(text.strip()) < 5:
            analysis["blank_pages"].append(i)

        content_hash = re.sub(r"\s+", "", text.lower())[:500]
        if content_hash:
            page_hashes.setdefault(content_hash, []).append(i)

        try:
            box = page.mediabox
            if float(box.width) > float(box.height):
                analysis["orientation_issues"].append(i)
        except Exception:
            pass

    for pages in page_hashes.values():
        if len(pages) > 1:
            analysis["duplicate_pages"].append(pages)

    if analysis["blank_pages"]:
        analysis["recommendations"].append(
            f"Remove {len(analysis['blank_pages'])} blank page(s) to save cost."
        )
    if analysis["duplicate_pages"]:
        analysis["recommendations"].append("Duplicate pages detected — consider removing duplicates.")
    if analysis["orientation_issues"]:
        analysis["recommendations"].append(
            f"{len(analysis['orientation_issues'])} page(s) may need landscape orientation."
        )

    return analysis
